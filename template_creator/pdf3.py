import tkinter as tk
from tkinter import messagebox
from tkinterPdfViewer import tkinterPdfViewer as pdf
import pymupdf
from PIL import Image, ImageTk
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import load_workbook
import shutil



class InvoiceSelector:

    def __init__(self):
        # self.pdf_file = pdf_file
        self.start_x = None
        self.start_y = None
        self.rect_id = None
        self.excel_data = {}
        self.selected_text = None
        self.location_data = None
        self.success = False
        self.fields = ["Supplier","Invoice_number","Date","Total"]
        self.page = None
        self.doc = None


    # get starting values of the pdf rectangle
    def start_selection(self, event):
        print(f"Mouse start to: {event.x}, {event.y}")
        self.start_x = event.x
        self.start_y = event.y 
    #get the end values of the pdf rectangle and pass them allong
    def end_selection(self, event):
        left = min(self.start_x,event.x)
        right= max(self.start_x,event.x)
        top = min(self.start_y,event.y)
        bottom = max(self.start_y,event.y)
        self.location_data = ", ".join([str(value) for value in [left, top, right, bottom]])
        rect = pymupdf.Rect(left, top, right, bottom)
        self.pdf_text_coordinate_value = self.page.get_text("text", clip=rect).strip()#text in the rectangle
        self.canvas.itemconfig(self.rect_id, outline="orange")
        print()
        print(f"Mouse end to: {event.x}, {event.y}")
    
    
    #allow the rectangle to be dragged
    def track_movement(self,event):
        # event.x and event.y give the coordinates relative to the widget
        left = min(self.start_x,event.x)
        right= max(self.start_x,event.x)
        top = min(self.start_y,event.y)
        bottom = max(self.start_y,event.y)
        if self.rect_id is not None:
            self.canvas.delete(self.rect_id)

        self.rect_id = self.canvas.create_rectangle(
            left, top, right, bottom, outline="black", width=3)

    def image_converter(self):
        pix = self.page.get_pixmap() 
        img = Image.frombytes("RGB",[pix.width, pix.height],pix.samples)
        photo = ImageTk.PhotoImage(img)
        return photo,pix.width, pix.height


    def validate(self,entry_value,entry_value_name):# used to confirm the values typed and selected from the pdf arethe same

        name = entry_value.get()
        if self.pdf_text_coordinate_value is None:
            messagebox.showwarning(
                "No Region Selected",
                "Please select a region from the PDF before validating.")
        else:
            if name.lower().strip() == self.pdf_text_coordinate_value.lower().strip():
                value = entry_value_name.removeprefix("entry_")
                self.excel_data[value] = [name,self.location_data]
                messagebox.showinfo(
                    "Validation Successful",
                    f"{entry_value_name} has been validated successfully."
                    )
                self.canvas.itemconfig(self.rect_id, outline="red")
                self.rect_id = None
                self.pdf_text_coordinate_value = None
                self.location_data = None
            else:
                messagebox.showwarning(
                    "Validation Failed",
                    f"Entered value: {name}\n"
                    f"PDF value: {self.pdf_text_coordinate_value}"
                )

    def create_excel_table(self):
        headers = ['Pdf_file_name']
        for item in self.fields:
            headers.append(item)
            item += "_coordinates"
            headers.append(item)
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "supplier_template"
        ws.append(headers)
        #syles used in the headers

        red_fill = PatternFill(start_color="FFC7CE",end_color="FFC7CE",fill_type="solid")
        header_fill = PatternFill(start_color="1F4E78",end_color="1F4E78",fill_type="solid")
        header_font = Font(name="Arial",size=11,bold=True,color="FFFFFF")         # White Text)
        header_alignment = Alignment(horizontal="center",vertical="center")

        #applying the styles
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        wb.save("supplier_template.xlsx")

    def submit_to_excel(self,pdf_file_name):
        if not all(item in self.excel_data for item in self.fields):
            messagebox.showwarning(
                "Incomplete Invoice",
                "Please validate all invoice fields before submitting."
            )
        else:
            path = Path("supplier_template.xlsx")

            if not path.exists():#if .xlsx fill does not exidt create it
                self.create_excel_table()

            wb = load_workbook("supplier_template.xlsx")
            ws = wb['supplier_template']

            target_row = ws.max_row + 1 
            header_map = {cell.value: cell.column for cell in ws[1] if cell.value is not None}
            ws.cell(row = target_row,column=1,value = pdf_file_name)######################################### pdffile
            
            for key, value in self.excel_data.items(): 
                if key in header_map:#adding value and is coordinates
                    col_num = header_map[key]
                    ws.cell(row = target_row,column=col_num,value = value[0] )
                    ws.cell(row = target_row,column=col_num+1,value = str(value[1]) )
            wb.save("supplier_template.xlsx")
            messagebox.showinfo(
                "Invoice Saved",
                f"{pdf_file_name} has been successfully added to the Excel template."
            )
            self.success = True
            self.root.destroy()
                
    #creating the right and left sections
    def GUI(self,pdf_file_name):# make the Graphic user intergace
        self.excel_data = {}
        self.root = tk.Tk()
        left_frame = tk.Frame(self.root)
        right_frame = tk.Frame(self.root)

        left_frame.grid(row=0, column=0)
        right_frame.grid(row=0, column=1)

        #making the left side
        photo, pix_width,pix_height = self.image_converter()
        self.canvas = tk.Canvas(left_frame, width=pix_width, height=pix_height, bg ="gray")
        self.canvas.image = photo
        self.canvas.create_image(0, 0, anchor="nw", image=photo)
        self.canvas.pack()

        # # Bind the motion event to the main window
        self.canvas.bind('<ButtonPress-1>', self.start_selection)
        self.canvas.bind('<B1-Motion>', self.track_movement)
        self.canvas.bind('<ButtonRelease-1>', self.end_selection)
        
        
        entries = {}

        for index, field in enumerate(self.fields):
            label = tk.Label(right_frame, text=f"{field}:")
            label.grid(row=index, column=0, padx=5, pady=5, sticky="e")
            
            entry = tk.Entry(right_frame)
            entry.grid(row=index, column=1, padx=5, pady=5)
            
            entries[field] = entry

            button = tk.Button(
                right_frame,
                text="Validate",
                command=lambda current_entry=entry, name=field: self.validate(current_entry, name)
            )
            button.grid(row=index, column=2, padx=5, pady=5)
        
        button_submit = tk.Button(
            right_frame,
            text="submit data",
            command = lambda :self.submit_to_excel(pdf_file_name),
        )
        button_submit.grid(column=1,pady=20, padx=5)
        self.root.mainloop()

    #creation and access of input and processed folders
    def access_documents(self):
        
        src_dir = Path("input")
        src_dir.mkdir(parents=True, exist_ok=True)
        dest_dir = Path("processed")
        dest_dir.mkdir(parents=True, exist_ok=True)
        for pdf_file in src_dir.glob("*.pdf"):
            self.success = False
            self.doc = pymupdf.open(pdf_file)
            self.page = self.doc[0]
            self.GUI(pdf_file.name)
            self.doc.close() 
            if self.success:
                dest = dest_dir / pdf_file.name
                pdf_file.rename(dest)
            self.success = False

if __name__ == "__main__":

    app = InvoiceSelector()
    app.access_documents()
